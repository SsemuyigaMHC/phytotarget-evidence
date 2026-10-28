#!/usr/bin/env python3
"""Create Chapter 13 claim/evidence templates without inventing evidence."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data'/'claims'/'Chapter13_Evidence_Matrix_Template.xlsx'

wb=Workbook(); ws=wb.active; ws.title='Evidence_Matrix'
cols=['Claim ID', 'Chapter section', 'Study PMID', 'DOI', 'Compound name / preparation', 'Compound identifier / structure', 'Protein / mediator', 'Protein isoform / organism', 'Proposed relationship', 'Prediction task', 'Computational evidence', 'Model and data provenance', 'Train / test separation', 'Predicted endpoint and score', 'Evidence origin', 'Experiment / assay', 'Measured endpoint (raw)', 'Protein source / biological system', 'Dose / exposure / timing', 'Controls and comparators', 'Observed result', 'Negative / discordant evidence', 'Claim supported', 'Unresolved alternatives', 'Status: tested positive / tested negative / untested / inconclusive', 'Source locator (figure/table/page)', 'Independent replication?', 'Reviewer decision', 'Scientific limitations', 'Figure or table mapping']
ws.append(cols)
for c in ws[1]: c.font=Font(bold=True); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
for i,h in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width=min(max(len(h)+2,12),30)

ws2=wb.create_sheet('Literature_Discovery_Master')
cols2=['Literature Record ID', 'PMID', 'DOI', 'PMCID', 'Title', 'Authors', 'Journal', 'Year', 'Source Databases', 'Originating Search Streams', 'Originating Search Queries', 'Primary Literature Category', 'Final Literature Category', 'Evidence Tier', 'Mechanism Flags', 'Subsection Hints', 'Scientific Role Flags', 'Foundational Status', 'Classification Status', 'Full Text Status', 'Full Text Path', 'Supplementary Count', 'Keep for Chapter?', 'Notes']
ws2.append(cols2)
for c in ws2[1]: c.font=Font(bold=True); c.alignment=Alignment(wrap_text=True,vertical='top')
ws2.freeze_panes='A2'; ws2.auto_filter.ref=ws2.dimensions
for i,h in enumerate(cols2,1): ws2.column_dimensions[get_column_letter(i)].width=min(max(len(h)+2,12),32)

ws3=wb.create_sheet('Coding_Guide')
rows=[('Core principle', 'Claim as unit', 'A compound–target–system–claim is the evidence unit; a paper may generate many distinct claims.'), ('Prediction', 'Predicted interaction is a hypothesis', 'Score, task, evaluation split and model provenance required; never upgrade to measured binding.'), ('Biochemical activity', 'Activity != binding', 'Inhibition, IC50 or activation under tested conditions do not alone establish Kd or engagement.'), ('Physical binding', 'Binding != intact-cell engagement', 'Record assay, protein source, Kd if measured, conditions, and controls.'), ('Cellular engagement', 'Preparation matters', 'CETSA in lysate is not automatically intact-cell engagement; stability shifts may be indirect.'), ('Functional contribution', 'Causality conditional', 'Perturbations/rescue, baseline effects and target function for binding-resistant variants.'), ('Evidence origin', 'Do not double count', 'Retrospective annotation; new within-study follow-up; independently replicated test remain distinct.'), ('Negative and missing', 'Do not collapse statuses', 'Untested, non-detectable under tested conditions, inconclusive, and undocumented are different.'), ('Computational model', 'Evaluation task matters', 'Warm versus cold; chemical novelty; prevalence; independent test; shortcut controls.'), ('Literature categories', 'A–F', 'See README: category is provisional scientific role, NOT an assigned evidence strength.'), ('Review design', 'Narrative purposive', 'Search and screening are targeted, not a systematic-review completeness claim.')]
for row in rows: ws3.append(row)
for c in ws3[1]: c.font=Font(bold=True)
ws3.column_dimensions['A'].width=24; ws3.column_dimensions['B'].width=60; ws3.column_dimensions['C'].width=90

OUT.parent.mkdir(parents=True,exist_ok=True); wb.save(OUT); print(OUT)
